import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel():
    with open('temp_seed_data.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="A30B1A", end_color="A30B1A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    regular_font = Font(name="Calibri", size=10, color="1F2937")
    bold_font = Font(name="Calibri", size=10, bold=True, color="111827")
    code_font = Font(name="Consolas", size=9.5, color="111827")

    # ----------------------------------------------------
    # 1. SHEET 1: ROUND 1 MCQs
    # ----------------------------------------------------
    ws_mcq = wb.create_sheet(title="Round 1 - MCQs (100)")
    ws_mcq.views.sheetView[0].showGridLines = True

    mcq_headers = [
        "Q No.", "Question ID", "Category", "Difficulty", "Question Text",
        "Option A", "Option B", "Option C", "Option D",
        "Correct Key", "Correct Answer Text", "Explanation"
    ]
    ws_mcq.append(mcq_headers)

    letters = ["A", "B", "C", "D"]

    for idx, q in enumerate(data.get("questions", []), start=1):
        opts = q.get("options", ["", "", "", ""])
        correct_idx = q.get("correctOption", 0)
        correct_letter = letters[correct_idx] if 0 <= correct_idx < len(letters) else str(correct_idx)
        correct_text = opts[correct_idx] if 0 <= correct_idx < len(opts) else ""

        row = [
            idx,
            q.get("questionId", f"MCQ-{idx:03d}"),
            q.get("category", ""),
            q.get("difficulty", "MEDIUM"),
            q.get("questionText", ""),
            opts[0] if len(opts) > 0 else "",
            opts[1] if len(opts) > 1 else "",
            opts[2] if len(opts) > 2 else "",
            opts[3] if len(opts) > 3 else "",
            correct_letter,
            correct_text,
            q.get("explanation", "")
        ]
        ws_mcq.append(row)

    # ----------------------------------------------------
    # 2. SHEET 2: ROUND 2 DEBUGGING PROBLEMS
    # ----------------------------------------------------
    ws_dbg = wb.create_sheet(title="Round 2 - Debugging (50)")
    ws_dbg.views.sheetView[0].showGridLines = True

    dbg_headers = [
        "Problem ID", "Title", "Language", "Difficulty", "Marks",
        "Description", "Buggy Code (Broken)", "Expected Output", "Fixed Solution Code", "Status"
    ]
    ws_dbg.append(dbg_headers)

    for p in data.get("debug", []):
        row = [
            p.get("problemId", ""),
            p.get("title", ""),
            p.get("language", ""),
            p.get("difficulty", ""),
            p.get("marks", 10),
            p.get("description", ""),
            p.get("brokenCode", ""),
            p.get("expectedOutput", ""),
            p.get("solutionSnippet", ""),
            p.get("status", "ACTIVE")
        ]
        ws_dbg.append(row)

    # ----------------------------------------------------
    # 3. SHEET 3: ROUND 3 TECH HUNT CLUES
    # ----------------------------------------------------
    ws_hunt = wb.create_sheet(title="Round 3 - Tech Hunt (50)")
    ws_hunt.views.sheetView[0].showGridLines = True

    hunt_headers = [
        "Station #", "Clue ID", "Title", "Category", "Clue Riddle / Challenge Text",
        "Secret Answer Key", "Roadmap Hint", "Hint Penalty", "Marks", "Status"
    ]
    ws_hunt.append(hunt_headers)

    for c in data.get("clues", []):
        row = [
            c.get("station", c.get("id", "")),
            c.get("clueId", ""),
            c.get("title", ""),
            c.get("category", ""),
            c.get("clueText", ""),
            c.get("answer", ""),
            c.get("hint", ""),
            c.get("hintPenalty", 2),
            c.get("marks", 10),
            c.get("status", "ACTIVE")
        ]
        ws_hunt.append(row)

    # ----------------------------------------------------
    # 4. SHEET 4: OFFICIAL USER ACCOUNTS
    # ----------------------------------------------------
    ws_users = wb.create_sheet(title="Official Accounts (Users)")
    ws_users.views.sheetView[0].showGridLines = True

    user_headers = [
        "User ID", "Name", "Email Address", "Role", "Assigned Round / Lab",
        "Password", "Security PIN", "Account Status", "Permissions"
    ]
    ws_users.append(user_headers)

    for u in data.get("users", []):
        row = [
            u.get("id", ""),
            u.get("name", ""),
            u.get("email", ""),
            u.get("role", ""),
            f"{u.get('department', '')} - {u.get('assignedRound', 'ALL')}",
            u.get("password", ""),
            u.get("pin", "-"),
            u.get("accountStatus", "ACTIVE"),
            ", ".join(u.get("permissions", []))
        ]
        ws_users.append(row)

    # ----------------------------------------------------
    # APPLY STYLING ACROSS ALL SHEETS
    # ----------------------------------------------------
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        # Format Header Row
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Format Data Rows
        for row_idx in range(2, ws.max_row + 1):
            is_even = (row_idx % 2 == 0)
            fill_to_use = zebra_fill if is_even else white_fill
            ws.row_dimensions[row_idx].height = 24

            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = fill_to_use
                cell.border = thin_border
                cell.font = regular_font

                # Code formatting for code columns in Sheet 2
                if ws.title.startswith("Round 2") and col_idx in [7, 8, 9]:
                    cell.font = code_font
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                # Text wrapping for long text columns
                elif any(word in str(ws.cell(row=1, column=col_idx).value).lower() for word in ['text', 'question', 'explanation', 'description', 'hint', 'code', 'permissions']):
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                elif col_idx in [1, 2, 4, 10] or 'id' in str(ws.cell(row=1, column=col_idx).value).lower() or 'key' in str(ws.cell(row=1, column=col_idx).value).lower():
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            header_val = str(col[0].value or '')
            
            for cell in col:
                val_str = str(cell.value or '')
                # If multi-line, check longest line
                lines = val_str.split('\n')
                longest = max(len(l) for l in lines) if lines else len(val_str)
                if longest > max_len:
                    max_len = longest

            # Column width bounds
            if 'text' in header_val.lower() or 'question' in header_val.lower() or 'code' in header_val.lower():
                width = min(max(max_len, 25), 60)
            elif 'explanation' in header_val.lower() or 'description' in header_val.lower():
                width = min(max(max_len, 25), 55)
            elif 'title' in header_val.lower() or 'answer' in header_val.lower():
                width = min(max(max_len, 18), 35)
            else:
                width = min(max(max_len + 4, 12), 30)
            
            ws.column_dimensions[col_letter].width = width

    # Save outputs
    output_filename_1 = "Festronix_TechNova_2026_Questions_Bank.xlsx"
    output_filename_2 = "Technova_Questions_Bank.xlsx"
    wb.save(output_filename_1)
    wb.save(output_filename_2)
    print(f"Successfully generated Excel files: {output_filename_1} and {output_filename_2}")

if __name__ == "__main__":
    build_excel()
